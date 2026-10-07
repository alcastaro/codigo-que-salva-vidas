"""Construye el deck «Código que salva vidas» con el sistema de diseño DevFest Santo Domingo.

Los fondos (tarjetas con muescas, badge, logos, gráficos) se componen en SVG a partir de los
archivos oficiales y se rasterizan con Chrome; todo el texto, las píldoras y las cajas son
objetos nativos y editables de PowerPoint.

    .venv/bin/python tools/construir_deck.py
"""
from __future__ import annotations

import base64
import re
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

RAIZ = Path(__file__).resolve().parents[1]
BRAND = RAIZ / "deck" / "brand" / "svg"
FIG = RAIZ / "deck" / "figuras"
CAPT = RAIZ / "deck" / "capturas"
# Enlace del evento de Climate Engine en el GDG Santo Domingo. PROVISIONAL: página del capítulo.
URL_EVENTO_CE = "https://gdg.community.dev/gdg-santo-domingo/"
FONDOS = RAIZ / "deck" / "fondos"
FONDOS.mkdir(parents=True, exist_ok=True)
SALIDA = RAIZ / "deck" / "codigo-que-salva-vidas.pptx"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

NEGRO, BLANCO, OFF = "#1e1e1e", "#ffffff", "#f0f0f0"
CORE = {"azul": "#4285f4", "verde": "#34a853", "amarillo": "#f9ab00", "rojo": "#ea4335"}
HALF = {"azul": "#57caff", "verde": "#5cdb6d", "amarillo": "#ffd427", "rojo": "#ff7daf"}
PASTEL = {"azul": "#c3ecf6", "verde": "#ccf6c5", "amarillo": "#ffe7a5", "rojo": "#f8d8d8"}
ICONO = {"azul": "cloud", "verde": "globe", "amarillo": "puzzle-hash", "rojo": "asterisk"}

SANS, SANS_MED = "Google Sans", "Google Sans Medium"
MONO_LIGHT, MONO_THIN = "Roboto Mono Light", "Roboto Mono Thin"
PX = 9525  # EMU por píxel en el lienzo de 960×540


# ============================================================================ fondos SVG
def _interior(nombre: str) -> tuple[str, str]:
    s = (BRAND / f"{nombre}.svg").read_text()
    vb = re.search(r'viewBox="([^"]+)"', s).group(1)
    cuerpo = re.sub(r"^.*?<svg[^>]*>", "", s, flags=re.S).rsplit("</svg>", 1)[0]
    return vb, cuerpo


def grafico(nombre: str, dx: float = 0, dy: float = 0, reemplazos: dict | None = None) -> str:
    """Un gráfico oficial en su posición nativa (coordenadas en pt → px ×4/3), desplazado dx, dy px."""
    _, cuerpo = _interior(nombre)
    for a, b in (reemplazos or {}).items():
        cuerpo = cuerpo.replace(a, b)
    return f'<g transform="translate({dx},{dy}) scale({4/3})">{cuerpo}</g>'


def logo(nombre: str, x: float, y: float, ancho: float) -> str:
    vb, cuerpo = _interior(nombre)
    vx, vy, vw, vh = map(float, vb.split())
    k = ancho / vw
    return f'<g transform="translate({x},{y}) scale({k}) translate({-vx},{-vy})">{cuerpo}</g>'


def ruta_redondeada(pts, r=28) -> str:
    """Polígono ortogonal con esquinas redondeadas (convexas y cóncavas) de radio r."""
    k = 0.5523
    n = len(pts)
    seg = []
    for i in range(n):
        p0, v, p2 = pts[i - 1], pts[i], pts[(i + 1) % n]

        def hacia(a, b, d):
            dx, dy = b[0] - a[0], b[1] - a[1]
            L = (dx * dx + dy * dy) ** 0.5
            return (a[0] + dx / L * d, a[1] + dy / L * d)
        a, b = hacia(v, p0, r), hacia(v, p2, r)
        c1 = (a[0] + (v[0] - a[0]) * k, a[1] + (v[1] - a[1]) * k)
        c2 = (b[0] + (v[0] - b[0]) * k, b[1] + (v[1] - b[1]) * k)
        seg.append((a, c1, c2, b))
    d = f"M{seg[0][0][0]:.2f},{seg[0][0][1]:.2f}"
    for i, (a, c1, c2, b) in enumerate(seg):
        if i:
            d += f" L{a[0]:.2f},{a[1]:.2f}"
        d += f" C{c1[0]:.2f},{c1[1]:.2f} {c2[0]:.2f},{c2[1]:.2f} {b[0]:.2f},{b[1]:.2f}"
    return d + " Z"


def tarjeta(pts, relleno=BLANCO) -> str:
    return (f'<path d="{ruta_redondeada(pts)}" fill="{relleno}" stroke="{NEGRO}" '
            f'stroke-width="2" stroke-linejoin="round"/>')


def imagen_recortada(pts, ruta_img: Path, id_: str) -> str:
    datos = base64.b64encode(ruta_img.read_bytes()).decode()
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    x, y, w, h = min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys)
    return (f'<clipPath id="{id_}"><path d="{ruta_redondeada(pts)}"/></clipPath>'
            f'<image href="data:image/png;base64,{datos}" x="{x}" y="{y}" width="{w}" height="{h}" '
            f'preserveAspectRatio="xMidYMid slice" clip-path="url(#{id_})"/>'
            f'<path d="{ruta_redondeada(pts)}" fill="none" stroke="{NEGRO}" stroke-width="2"/>')


LOGO_NOTCH = lambda: logo("logo-gdg-stacked", 716, 450, 201)  # noqa: E731
CARD = [(38, 40), (922, 40), (922, 422), (679, 422), (679, 501), (38, 501)]
TITULO = [(336, 40), (922, 40), (922, 422), (679, 422), (679, 501), (38, 501), (38, 175), (336, 175)]
SECCION = [(38, 40), (922, 40), (922, 422), (679, 422), (679, 501), (288, 501), (288, 156), (38, 156)]
IZQ = [(38, 40), (461, 40), (461, 501), (38, 501)]
FOTO = [(499, 40), (922, 40), (922, 422), (679, 422), (679, 501), (499, 501)]
BADGE = {'font-size="26">Track': 'font-size="21.5">Workshops', "Editable Location": "Laboratorio"}


def svg_a_png(nombre: str, contenido: str, fondo=OFF) -> Path:
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 960 540" width="1920" height="1080">'
           f'<rect width="960" height="540" fill="{fondo}"/>{contenido}</svg>')
    html = FONDOS / f"{nombre}.html"
    html.write_text(f"<html><body style='margin:0'>{svg}</body></html>")
    png = FONDOS / f"{nombre}.png"
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--screenshot={png}",
                    "--window-size=1920,1080", "--force-device-scale-factor=1", html.as_uri()],
                   check=True, capture_output=True)
    html.unlink()
    return png


def fondos() -> dict[str, Path]:
    f = {}
    f["titulo"] = svg_a_png("titulo", tarjeta(TITULO) + grafico("track-badge", 8, 2, BADGE) + LOGO_NOTCH())
    f["card"] = svg_a_png("card", tarjeta(CARD) + LOGO_NOTCH())
    f["seccion"] = svg_a_png("seccion", tarjeta(SECCION) + grafico("arrow-up") + grafico("scallop-arches") + LOGO_NOTCH())
    f["chart"] = svg_a_png("chart", "")
    puntos = "".join(f'<circle cx="{714 + i * 84}" cy="455" r="32" fill="{CORE["azul"]}" stroke="{NEGRO}" stroke-width="2"/>' for i in range(3))
    f["statement"] = svg_a_png(
        "statement",
        grafico("scallop-bracket-left") + grafico("scallop-bracket-right", 0, 0)
        + grafico("track-badge", 622, 0, {**BADGE, PASTEL["azul"]: CORE["azul"]})
        + grafico("cloud", 620, 50) + puntos, fondo=PASTEL["azul"])
    f["numero"] = svg_a_png(
        "numero",
        grafico("scallop-bracket-short-left") + grafico("scallop-bracket-short-right")
        + grafico("double-slash-wide")
        + f'<circle cx="591" cy="92" r="58" fill="{PASTEL["amarillo"]}" stroke="{NEGRO}" stroke-width="2"/>'
        + f'<circle cx="872" cy="447" r="58" fill="{PASTEL["rojo"]}" stroke="{NEGRO}" stroke-width="2"/>')
    puntos_v = "".join(f'<circle cx="{714 + i * 84}" cy="455" r="32" fill="{CORE["verde"]}" stroke="{NEGRO}" stroke-width="2"/>' for i in range(3))
    f["statement_verde"] = svg_a_png(
        "statement_verde",
        grafico("scallop-bracket-left") + grafico("scallop-bracket-right", 0, 0)
        + grafico("track-badge", 622, 0, {'font-size="26">Track': 'font-size="21.5">Comunidad', "Editable Location": "Santo Domingo",
                                          PASTEL["azul"]: CORE["verde"]})
        + puntos_v, fondo=PASTEL["verde"])
    for nombre, img in [("foto_melissa", "melissa_ensamble.png"), ("foto_alphaearth", "alphaearth_cuadro.png"),
                        ("foto_radar", "radar_cuadro.png")]:
        if (FIG / img).exists():
            f[nombre] = svg_a_png(nombre, tarjeta(IZQ) + imagen_recortada(FOTO, FIG / img, nombre) + LOGO_NOTCH())
    return f


# ============================================================================ pptx
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def caja_texto(sl, x, y, w, h, parrafos, fuente=SANS, px=33, color=NEGRO, interlineado=1.0,
               alinear=PP_ALIGN.LEFT, ancla=MSO_ANCHOR.TOP, antes_px=0):
    """parrafos: lista de str o de listas de (texto, negrita) para mezclar pesos."""
    tb = sl.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancla
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        p.line_spacing = interlineado
        if antes_px and i:
            p.space_before = Pt(antes_px * 0.75)
        trozos = [(par, False)] if isinstance(par, str) else par
        for texto, negrita in trozos:
            r = p.add_run()
            r.text = texto
            r.font.name = fuente
            r.font.size = Pt(px * 0.75)
            r.font.bold = negrita
            r.font.color.rgb = rgb(color)
    return tb


def viñetas(sl, x, y, w, h, items, px=33):
    tb = caja_texto(sl, x, y, w, h, items, px=px, interlineado=1.15, antes_px=13)
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(34 * PX)))
        pPr.set("indent", str(int(-34 * PX)))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
    return tb


def forma(sl, tipo, x, y, w, h, relleno=BLANCO, radio_px=None, texto=None, px=19, fuente=SANS_MED,
          color_texto=NEGRO, alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE):
    s = sl.shapes.add_shape(tipo, Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(relleno)
    s.line.color.rgb = rgb(NEGRO)
    s.line.width = Pt(1.5)
    s.shadow.inherit = False
    if tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.5 if radio_px is None else min(0.5, radio_px / min(w, h))
    if texto is not None:
        tf = s.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Emu(10 * PX)
        tf.margin_top = tf.margin_bottom = Emu(4 * PX)
        tf.vertical_anchor = ancla
        for i, linea in enumerate(texto.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = alinear
            r = p.add_run()
            r.text = linea
            r.font.name, r.font.size = fuente, Pt(px * 0.75)
            r.font.color.rgb = rgb(color_texto)
    return s


def pildora(sl, x, y, w, h, texto, relleno=BLANCO, px=33, fuente=SANS, color_texto=NEGRO):
    return forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, relleno, None, texto, px, fuente, color_texto)


def flecha(sl, x1, y1, x2, y2, color=CORE["azul"]):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1 * PX), Emu(y1 * PX), Emu(x2 * PX), Emu(y2 * PX))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


class Deck:
    def __init__(self, fondos: dict[str, Path]):
        self.f = fondos
        self.p = Presentation()
        self.p.slide_width, self.p.slide_height = Emu(960 * PX), Emu(540 * PX)
        self.blank = self.p.slide_layouts[6]

    def nueva(self, fondo: str, notas: str = ""):
        sl = self.p.slides.add_slide(self.blank)
        sl.shapes.add_picture(str(self.f[fondo]), 0, 0, self.p.slide_width, self.p.slide_height)
        if notas:
            sl.notes_slide.notes_text_frame.text = notas
        return sl

    # ---- diseños
    def titulo(self, lineas, ponente, notas=""):
        sl = self.nueva("titulo", notas)
        caja_texto(sl, 86, 196, 760, 190, lineas, fuente=SANS, px=80, interlineado=0.9,
                   ancla=MSO_ANCHOR.BOTTOM).text_frame.paragraphs  # noqa: B018
        for p in sl.shapes[-1].text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
        caja_texto(sl, 86, 404, 560, 60, ponente, fuente=MONO_LIGHT, px=21, interlineado=1.0)

    def seccion(self, numero, titulo, sub, notas=""):
        sl = self.nueva("seccion", notas)
        caja_texto(sl, 68, 58, 200, 90, [numero], fuente=MONO_THIN, px=81)
        t = caja_texto(sl, 344, 150, 540, 200, titulo, px=60, interlineado=0.9, ancla=MSO_ANCHOR.BOTTOM)
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
        caja_texto(sl, 344, 368, 320, 110, sub, fuente=MONO_LIGHT, px=24, interlineado=1.15)

    def vinetas(self, titulo, items, notas="", sub=None):
        sl = self.nueva("card", notas)
        t = caja_texto(sl, 90, 82, 780, 56, [titulo], px=40)
        t.text_frame.paragraphs[0].runs[0].font.bold = True
        y = 150
        if sub:
            caja_texto(sl, 90, 140, 780, 30, [sub], fuente=MONO_LIGHT, px=21)
            y = 182
        viñetas(sl, 96, y, 760, 260, items)

    def foto(self, fondo, titulo, items, notas="", credito=None):
        sl = self.nueva(fondo, notas)
        t = caja_texto(sl, 82, 82, 340, 100, titulo, px=40)
        for p in t.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
        y = 82 + 50 * len(titulo) + 18
        viñetas(sl, 86, y, 350, 470 - y, items, px=27)
        if credito:
            caja_texto(sl, 499, 508, 420, 24, [credito], fuente=MONO_LIGHT, px=11)

    def frase(self, lider, resto, notas=""):
        sl = self.nueva("statement", notas)
        caja_texto(sl, 152, 120, 400, 300, [[(lider, True), (resto, False)]], px=33, interlineado=1.15,
                   ancla=MSO_ANCHOR.MIDDLE)

    def numero(self, arriba, cifra, abajo, notas=""):
        sl = self.nueva("numero", notas)
        pildora(sl, 32, 34, 472, 72, arriba)
        caja_texto(sl, 60, 140, 416, 260, [cifra], fuente=SANS_MED, px=147, alinear=PP_ALIGN.CENTER,
                   ancla=MSO_ANCHOR.MIDDLE)
        pildora(sl, 32, 434, 472, 72, abajo)

    def portafolio(self, titulo, columnas, notas="", ancho_util=884, extra=None):
        """columnas: lista de (etiqueta, color, [cajas])."""
        sl = self.nueva("chart", notas)
        caja_texto(sl, 37, 30, 880, 50, [titulo], px=33)
        n = len(columnas)
        hueco = 22
        w = (ancho_util - hueco * (n - 1)) / n
        for i, (etq, col, cajas) in enumerate(columnas):
            x = 35 + i * (w + hueco)
            pildora(sl, x, 112, w, 42, etq, relleno=HALF[col], px=19, fuente=SANS_MED)
            for j, c in enumerate(cajas):
                forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, 178 + j * 108, w, 88, PASTEL[col], 12, c, px=19)
        if extra:
            extra(sl)
        return sl

    def flujo(self, titulo, pasos, pie=None, notas=""):
        sl = self.nueva("chart", notas)
        caja_texto(sl, 37, 30, 880, 50, [titulo], px=33)
        n, hueco = len(pasos), 34
        w = (888 - hueco * (n - 1)) / n
        cols = ["verde", "azul", "azul", "azul", "rojo"]
        for i, (etq, texto) in enumerate(pasos):
            x = 36 + i * (w + hueco)
            col = cols[i] if i < len(cols) else "azul"
            pildora(sl, x, 150, w, 42, etq, relleno=HALF[col], px=19, fuente=SANS_MED)
            forma(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, 214, w, 150, PASTEL[col], 12, texto, px=19)
            if i < n - 1:
                flecha(sl, x + w + 4, 289, x + w + hueco - 4, 289, CORE["azul"])
        if pie:
            caja_texto(sl, 37, 410, 880, 80, pie, fuente=MONO_LIGHT, px=21, interlineado=1.15)
        return sl

    def pantalla(self, titulo, captura, puntos, credito, notas="", ancho_img=620):
        """Captura o GIF de una plataforma (esquinas de 18 px, contorno de 2 px) con puntos clave a la derecha."""
        sl = self.nueva("chart", notas)
        caja_texto(sl, 37, 30, 880, 50, [titulo], px=33)
        ruta = CAPT / captura
        from PIL import Image
        with Image.open(ruta) as im:
            w0, h0 = im.size
        w = ancho_img
        h = w * h0 / w0
        if h > 390:
            h, w = 390, 390 * w0 / h0
        pic = sl.shapes.add_picture(str(ruta), Emu(37 * PX), Emu(104 * PX), Emu(int(w * PX)), Emu(int(h * PX)))
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = geom.makeelement(qn("a:avLst"), {})
            geom.append(av)
        av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": f"val {int(18 / min(w, h) * 100000)}"}))
        pic.line.color.rgb = rgb(NEGRO)
        pic.line.width = Pt(1.5)
        x = 37 + w + 28
        viñetas(sl, x, 110, 923 - x, 330, puntos, px=21)
        caja_texto(sl, 37, 104 + h + 10, w, 24, [credito], fuente=MONO_LIGHT, px=13)
        return sl

    def evento(self, lider, resto, qr_png, url_corta, notas=""):
        sl = self.nueva("statement_verde", notas)
        caja_texto(sl, 152, 110, 400, 320, [[(lider, True), (resto, False)]], px=33, interlineado=1.15,
                   ancla=MSO_ANCHOR.MIDDLE)
        sl.shapes.add_picture(str(qr_png), Emu(690 * PX), Emu(160 * PX), Emu(200 * PX), Emu(200 * PX))
        caja_texto(sl, 640, 368, 300, 40, [url_corta], fuente=MONO_LIGHT, px=14, alinear=PP_ALIGN.CENTER)

    def qr_grande(self, titulo, qr_png, url, notas=""):
        sl = self.nueva("chart", notas)
        caja_texto(sl, 37, 30, 880, 50, [titulo], px=33)
        sl.shapes.add_picture(str(qr_png), Emu(340 * PX), Emu(110 * PX), Emu(280 * PX), Emu(280 * PX))
        caja_texto(sl, 180, 410, 600, 40, [url], fuente=MONO_LIGHT, px=21, alinear=PP_ALIGN.CENTER)

    def guardar(self):
        self.p.save(SALIDA)
        print("escrito", SALIDA.relative_to(RAIZ), "·", len(self.p.slides), "diapositivas")


# ============================================================================ contenido
def qr_png(url: str, nombre: str) -> Path:
    import qrcode
    q = qrcode.QRCode(border=1, box_size=20, error_correction=qrcode.constants.ERROR_CORRECT_M)
    q.add_data(url)
    q.make(fit=True)
    ruta = FIG / f"qr_{nombre}.png"
    q.make_image(fill_color=NEGRO, back_color=BLANCO).save(str(ruta))
    return ruta


def construir():
    d = Deck(fondos())
    PORTAL = "alcastaro.github.io/codigo-que-salva-vidas"

    # ------------------------------------------------------------------ apertura
    d.titulo(["Código que", "salva vidas"],
             ["Alberto Castillo Aroca", "Tech Lead en contextos de crisis"],
             "Presentación (30 s). Organizador del GDG Santo Domingo, Google.org Gen AI Accelerator. "
             "Hoy: qué construye Google para crisis y, sobre todo, lo tocamos con las manos. "
             "Pide que escaneen el QR desde ya: el notebook tarda ~30 s en instalar.")

    d.frase("En una crisis", " no faltan datos: sobra ruido, y lo que se sabe no llega a quien lo necesita.",
            "El problema no es la falta de datos sino la asimetría y el ruido. Todo lo que veremos son "
            "piezas de una misma cadena: detectar → predecir → avisar a tiempo y de forma comprensible.")

    d.numero("Flood Hub pronostica en", "150+", "países con pronóstico",
             "Fuente: crisisresilience.google — pronósticos de inundación para 2.000 millones de personas en "
             "más de 150 países. Otras cifras de la misma página: alertas sísmicas en 95+ países, alianzas de "
             "alertas en 90+ países, 10 millones de conexiones diarias con información de crisis (promedio 2025).")

    d.pantalla("crisisresilience.google", "crisis.gif", [
        "Incendios", "Inundaciones", "Sismos", "Ciclones y clima", "Alertas en 90+ países",
    ], "crisisresilience.google · captura oct-2026",
        "El portal que reúne todo. 'Nadie debería ser sorprendido por un desastre natural'. Recorrer rápido "
        "(20 s): los cinco frentes y las historias de socios (GiveDirectly en Nigeria, WeatherNext y Melissa).")

    d.portafolio("Google Crisis Resilience: el mapa", [
        ("Sismos", "rojo", ["Android\nEarthquake Alerts", "Red de acelerómetros\nen teléfonos", "Alertas en\n98 países"]),
        ("Inundaciones", "azul", ["Flood Hub", "Flood Forecasting\nAPI", "150+ países,\nríos sin estaciones"]),
        ("Ciclones y clima", "amarillo", ["WeatherNext 2", "Weather Lab\n(ciclones)", "Hasta 15 días,\nprobabilístico"]),
        ("Incendios", "verde", ["Límites de incendio\nen Search y Maps", "FireSat\n(constelación)", "Detección de\n5×5 m en 20 min"]),
    ], "Cuatro amenazas, una lógica. Todos comparten capas: satélite, modelos de IA y un canal de alertas "
       "(SOS Alerts / Public Alerts en Search, Maps y Android). Fuentes: crisisresilience.google; research.google.")

    d.portafolio("Las capas que lo sostienen todo", [
        ("Observar", "verde", ["AlphaEarth\nFoundations", "Earth Engine\ny Google Earth", "Sentinel / Landsat\n(datos abiertos)"]),
        ("Predecir", "azul", ["Modelos LSTM\n(Flood Hub)", "WeatherNext 2", "Earth AI\n(Gemini)"]),
        ("Avisar", "rojo", ["SOS Alerts", "Public Alerts\n(estándar CAP)", "Android, Search\ny Maps"]),
        ("Construir", "amarillo", ["Gemini API\n(AI Studio)", "Gemma\n(modelos abiertos)", "Colab"]),
    ], "La vista para developers: qué puedo usar yo. En el taller tocamos Observar (Sentinel, AlphaEarth), "
       "Predecir (Weather Lab, Flood Hub) y Construir (Gemini, Gemma, Colab, AI Studio).")

    # ------------------------------------------------------------------ pilar 1
    d.seccion("01", ["Telemetría", "distribuida"], ["Millones de teléfonos", "funcionando como", "sismógrafos"],
              "Pilar 1. El acelerómetro que gira tu pantalla también siente un sismo.")

    d.flujo("Android Earthquake Alerts: cómo funciona", [
        ("Sensor", "El acelerómetro siente la sacudida"),
        ("Teléfono", "Un algoritmo local decide si es sismo"),
        ("Señal", "Envía solo hora y zona aproximada"),
        ("Servidor", "Une miles de señales: epicentro y magnitud"),
        ("Alerta", "Llega antes que la onda que más daña"),
    ], ["Funciona porque las ondas viajan a unos km/s y los datos a la velocidad de la luz.",
        "Fuente: Allen et al., Science (julio 2025)."],
        "Arquitectura: detección en el borde (on-device) + agregación en la nube. El teléfono no manda la "
        "señal cruda, solo un disparo con ubicación gruesa: privacidad y ancho de banda.")

    d.numero("recibió la alerta", "36 %", "antes de sentir la sacudida",
             "Science (2025): alertas en 98 países para M≥4.5, ~60 eventos y 18 millones de alertas al mes; 85 % "
             "de quienes recibieron alerta sintieron el sismo; 36 % la recibió antes, 28 % durante y 23 % después.")

    # ------------------------------------------------------------------ pilar 2
    d.seccion("02", ["Modelado", "predictivo"], ["Del satélite y la", "lluvia, al río y", "al huracán"],
              "Pilar 2. Aquí entra la IA de verdad: Flood Hub, WeatherNext, AlphaEarth y Earth Engine.")

    d.pantalla("Flood Hub: ríos que avisan", "floodhub.gif", [
        "Dos redes LSTM: pasado y pronóstico",
        "Ríos sin estaciones de medición",
        "Hasta 7 días antes",
        "API con lista de espera",
    ], "sites.research.google/floods · captura oct-2026",
        "Nature (2024), 'Global prediction of extreme floods in ungauged watersheds'. Un LSTM 'hindcast' lee un "
        "año de clima; otro LSTM 'forecast' toma ese estado y 7 días de pronóstico meteorológico. La API exige key "
        "aprobada y uso no comercial: en el taller uso mi key y comparto una copia.")

    d.pantalla("WeatherNext 2", "weathernext.gif", [
        "8× más rápido",
        "Mejor en 99,9 % de variables",
        "Hasta 15 días, por hora",
        "En Earth Engine, BigQuery y Vertex AI",
    ], "deepmind.google/science/weathernext · captura oct-2026",
        "WeatherNext 2 (Google DeepMind, nov-2025): 8× más rápido, supera a la versión anterior en 99,9 % de "
        "variables y plazos, hasta 15 días. Datos en Earth Engine y BigQuery; inferencia en Vertex AI.")

    d.pantalla("Melissa: 50 futuros posibles", "melissa_animado.gif", [
        "Weather Lab, emitido el 25-oct",
        "Azul: escenarios del modelo",
        "Rojo: lo que pasó (NHC)",
        "Esto lo haces en la Pista A",
    ], "Weather Lab (CC BY 4.0) · NOAA/NHC · Natural Earth · animación propia",
        "Weather Lab (deepmind.google.com/science/weatherlab) ahora pide iniciar sesión, así que la animación está "
        "hecha con sus CSV públicos. Cada línea azul es un escenario del modelo experimental de ciclones; la roja, "
        "el best track del NHC. Pausa aquí: '¿el abanico contiene la línea roja?'.")

    d.numero("Cinco días antes:", "80 %", "de confianza en categoría 5",
             "DeepMind (mayo 2026): WeatherNext predijo la llegada a Jamaica con fuerza de categoría 5 cinco días "
             "antes con 80 % de confianza, casi 100 % tres días antes. Mejor modelo individual de trayectoria e "
             "intensidad de 2025 según el NHC. Melissa tocó tierra el 28-oct-2025 con 160 nudos.")

    d.pantalla("Sentinel-1: el radar ve bajo las nubes", "radar_parpadeo.gif", [
        "Atraviesa nubes, de día y de noche",
        "Agua quieta = oscuro",
        "Antes 17-oct · después 29-oct",
        "Datos abiertos (Copernicus)",
    ], "Copernicus Sentinel-1 (ESA) · Microsoft Planetary Computer · animación propia",
        "No es de Google, pero es la materia prima. Durante Melissa las imágenes ópticas tenían 100 % de nubes. "
        "Black River (Jamaica), misma órbita antes y después; en celeste, agua nueva. Pista B.", ancho_img=420)

    d.foto("foto_alphaearth", ["AlphaEarth", "Foundations"], [
        "64 números por píxel",
        "Busca sin entrenar",
        "Anual, 2017–2025",
        "Abierto y gratuito",
    ], "AlphaEarth Foundations (DeepMind, jul-2025): embedding de 64 dimensiones por píxel de 10 m; más de 1,4 "
       "billones de huellas al año; 24 % menos error que los modelos comparados. En rojo, todo lo que 'se parece' "
       "al centro de Black River (similitud > 0,9): encuentra los asentamientos sin entrenar nada. Pistas B y D.",
       credito="AlphaEarth Foundations (CC BY 4.0) · Source Cooperative")

    if (CAPT / "alphaearth_constanza.gif").exists():
        d.pantalla("AlphaEarth: Constanza, 2017 → 2025", "alphaearth_constanza.gif", [
            "Un color por tipo de terreno",
            "Derecha: cuánto cambió",
            "Cultivos en laderas",
            "Esto lo haces en la Pista D",
        ], "AlphaEarth Foundations (CC BY 4.0) vía Source Cooperative · animación propia",
            "Misma proyección de color para todos los años, para que sean comparables. Constanza: agricultura "
            "intensiva en un valle de montaña. Las zonas claras a la derecha cambiaron respecto a 2017. Es una señal "
            "para revisar, no un diagnóstico de erosión.")

    d.pantalla("Earth Engine: 40 años de Santo Domingo", "timelapse_sdq.gif", [
        "Petabytes de satélite y clima",
        "Análisis en la nube, en minutos",
        "Gratis para uso no comercial",
        "Base de Timelapse y Climate Engine",
    ], "earthengine.google.com/timelapse · captura oct-2026",
        "Earth Engine es la plataforma: catálogo de datos + cómputo. Timelapse muestra 1984–2022. El catálogo incluye "
        "el Satellite Embedding de AlphaEarth. En el taller no lo usamos para no obligar a crear un proyecto de Cloud.")

    d.pantalla("Novedades de Google Earth", "googleearth.gif", [
        "Earth AI: imagen, población y entorno",
        "Gemini razona sobre el mapa",
        "Classify: cultivos con AlphaEarth",
        "Pro y AI Pro, primero en EE. UU.",
    ], "earth.google.com · captura oct-2026",
        "Google Earth AI (oct-2025): modelos de Imagery, Population y Environment + un agente de Geospatial Reasoning "
        "con Gemini (0,82 de exactitud frente a 0,50 de Gemini 2.5 Pro solo). En Google Earth: 'Ask Google Earth' "
        "con Gemini y Classify (experimental), que con ≥3 ejemplos por clase y AlphaEarth clasifica cada 10 m con un "
        "Random Forest. Disponibilidad inicial: Google Earth Professional y Google AI Pro en EE. UU.")

    d.pantalla("Amazon Mining Watch", "amw.gif", [
        "Minería de oro en 9 países",
        "Redes neuronales sobre Sentinel-2",
        "99 % especificidad, 90 % sensibilidad",
        "Datos abiertos, cada trimestre",
    ], "amazonminingwatch.org (Earth Genome y aliados) · captura oct-2026",
        "Earth Genome con Amazon Conservation, la Fundación Moore, Pulitzer Center y organizaciones indígenas. "
        "Actualización de agosto 2026: nuevo ensamble de CNN sobre Sentinel-2 en cuadros de 480 m, segmentación con "
        "SAM2, datos 2018–2026 T2; 1,149 millones de hectáreas afectadas. Datos en Source Cooperative. Es el modelo "
        "de la Pista D: vigilar cambios del suelo y priorizar dónde mirar.")

    d.pantalla("Climate Engine", "climateengine.png", [
        "Clima y satélite sin código",
        "Corre sobre Earth Engine",
        "Sequía, incendios y cultivos",
        "Web, API e informes",
    ], "climateengine.org · captura oct-2026",
        "Climate Engine: aplicación de Desert Research Institute y Universidad de Idaho con Google, sobre Earth "
        "Engine. Procesa en segundos 30+ años de Landsat y datos climáticos para sequía, uso de agua, agricultura e "
        "incendios; resultados descargables y compartibles por enlace. Puente al anuncio del evento del GDG.")

    # ------------------------------------------------------------------ pilar 3
    d.seccion("03", ["Despliegue", "a escala"], ["Cuando la", "infraestructura", "física colapsa"],
              "Pilar 3. La mejor predicción del mundo no sirve si no llega y no se entiende.")

    d.vinetas("La última milla: la alerta", [
        "SOS Alerts en Search y Maps",
        "Alertas oficiales con estándar CAP",
        "Límites de incendio cada 10–15 min",
        "Si no se entiende en 3 s, no sirve",
    ], "SOS Alerts reúne información de autoridades en Search y Maps. Public Alerts recibe alertas oficiales en "
       "CAP (Common Alerting Protocol): severidad, urgencia y certeza son campos del estándar. En el taller, Gemini "
       "redacta alertas con esos campos.")

    d.portafolio("La comunidad también construye: Gemma 4 Good", [
        ("Acuífero·Vigía", "azul", ["Alerta de inundación\nsin internet", "Gemma 4 E2B en\nRaspberry Pi 5", "Sudamérica ·\nGlobal Resilience"]),
        ("PreVillage", "verde", ["Servicios públicos\npor WhatsApp", "Gemma 4 E4B en\nRaspberry Pi 5", "Nepal ·\npremio llama.cpp"]),
        ("Gilbeot", "amarillo", ["Navegación con\npuntos de referencia", "Gemma 4 E2B\nen Android", "Personas mayores ·\npremio LiteRT"]),
    ], "Gemma 4 Good Challenge (Kaggle + Google DeepMind), ganadores anunciados el 24-ago-2026, 1.600+ proyectos. "
       "Lo que Google hace a escala planetaria, la comunidad lo lleva al borde, donde no hay señal.")

    # ------------------------------------------------------------------ taller
    d.seccion("04", ["Manos", "a la obra"], ["Elige tu pista:", "huracán, satélite,", "alerta o AlphaEarth"],
              "Taller (25 min). Todos abren el portal del QR. No hace falta proyecto de Google Cloud.")

    d.pantalla("Pista D: tu plataforma satelital", "plataforma_app.gif", [
        "Cambio 2017→2025 con AlphaEarth",
        "× pendiente Copernicus",
        "Clic para dar ejemplos",
        "Enlace público con Gradio",
    ], "Plataforma del taller (Gradio en Colab) · Constanza, RD",
        "La idea de Amazon Mining Watch y de Classify, en 25 minutos y para RD. Cambio fuerte en terreno empinado = "
        "señal para revisar, no diagnóstico. En Constanza el cambio está sobre todo en el valle (rotación de "
        "cultivos); las señales aparecen en el borde de las laderas. Código escrito desde cero para el taller.")

    d.portafolio("Elige tu pista", [
        ("A · Huracán", "azul", ["Weather Lab:\n50 escenarios", "Flood Hub:\nríos en alerta", "Gemini: la alerta"]),
        ("B · Satélite", "verde", ["Sentinel-1 y 2\ncon STAC", "Radar: agua\nbajo las nubes", "AlphaEarth:\nexposición"]),
        ("C · Alerta", "rojo", ["AI Studio,\nsin código", "Formato CAP\ny Gemma", "Build: tu app"]),
        ("D · AlphaEarth", "amarillo", ["Erosión en RD:\ncambio y pendiente", "Clasifica con\nejemplos", "Tu plataforma\n(Gradio)"]),
    ], "A, B y D: Google Colab (laptop). C: AI Studio, funciona desde el teléfono. Key de Gemini gratis en "
       "aistudio.google.com/apikey, guardada en Secretos de Colab (nunca pegada en una celda).")

    d.qr_grande("Escanea y elige tu pista", FIG / "qr_portal.png", PORTAL,
                "Dejar esta diapositiva en pantalla durante el taller.")

    d.evento("Próximamente en el GDG Santo Domingo:",
             " un taller de Climate Engine para analizar clima y satélite sin código.",
             qr_png(URL_EVENTO_CE, "evento_ce"), URL_EVENTO_CE.replace("https://", "").rstrip("/"),
             "Anuncio del evento de Climate Engine del GDG Santo Domingo. ⚠️ El QR apunta de forma provisional a la "
             "página del capítulo; cambiar URL_EVENTO_CE en tools/construir_deck.py por el enlace del evento.")

    d.vinetas("Llévatelo a casa", [
        "Notebooks y guía en el portal del QR",
        "Datos abiertos, con su licencia citada",
        "Alertas reales: tu servicio meteorológico",
        "¿Construyes algo? Cuéntaselo al GDG",
    ], "Cierre. Lo que vimos es investigación y práctica: las alertas reales las emite una autoridad.",
        sub="Gracias · Código que salva vidas · DevFest Santo Domingo 2026")

    d.guardar()


if __name__ == "__main__":
    construir()
